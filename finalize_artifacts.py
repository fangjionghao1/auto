from pathlib import Path
import re, zipfile
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT=Path(__file__).parent
for paper in sorted((ROOT/'gen'/'artifacts').glob('*/*/*')):
    if not (paper/'论文.md').exists():
        continue
    sim=paper/'仿真'
    (sim/'env.txt').write_text('MATLAB R2025b; base MATLAB only; random seed 42; synthetic scenario.\n',encoding='utf-8')
    for doc in paper.glob('*.docx'):
        d=Document(doc)
        sec=d.sections[0]
        footer=sec.footer.paragraphs[0]
        footer.alignment=1
        footer.add_run('第 ')
        fld=OxmlElement('w:fldSimple'); fld.set(qn('w:instr'),'PAGE')
        footer._p.append(fld)
        footer.add_run(' 页')
        d.save(doc)
    with zipfile.ZipFile(sim/'仿真打包.zip','w',zipfile.ZIP_DEFLATED) as z:
        for path in sim.rglob('*'):
            if path.is_file() and path.name!='仿真打包.zip':
                z.write(path,path.relative_to(sim))
    body=(paper/'论文.md').read_text(encoding='utf-8')
    section=body.split('## 参考文献')[0].split('## References')[0]
    chars=len(re.findall(r'[\u4e00-\u9fff]',section))
    print(paper.name, 'Chinese chars:',chars, 'zip files:',len(zipfile.ZipFile(sim/'仿真打包.zip').namelist()))
